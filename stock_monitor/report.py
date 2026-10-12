from pathlib import Path
import subprocess
import tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def money(value, currency='USD'):
    if value is None:
        return 'N/A'
    for scale, suffix in [(1e12, 'T'), (1e9, 'B'), (1e6, 'M')]:
        if abs(value) >= scale:
            return f'{currency} {value/scale:,.2f}{suffix}'
    return f'{currency} {value:,.2f}'


def multiple(value):
    return 'N/M' if value is not None and value <= 0 else ('N/A' if value is None else f'{value:,.2f}x')


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string('17365D')
    p.paragraph_format.keep_with_next = True


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for row in t.rows:
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def hyperlink(paragraph, text, url):
    if not str(url).startswith(('https://', 'http://')):
        paragraph.add_run(text)
        return
    from docx.opc.constants import RELATIONSHIP_TYPE
    rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    r.append(props)
    t = OxmlElement('w:t'); t.text = text; r.append(t)
    link.append(r); paragraph._p.append(link)


def build_report(path, company, start, end, valuation, summary, metadata, sotp=None, font_size=10):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'; style.font.size = Pt(font_size)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.0
    footer = section.footer.paragraphs[0]
    footer.text = f'{company["ticker"]} | Research dossier | '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    title = doc.add_paragraph()
    r = title.add_run(f'{company["name"]} ({company["ticker"]})')
    r.bold = True; r.font.size = Pt(17); r.font.color.rgb = RGBColor.from_string('17365D')
    doc.add_paragraph('Institutional Intelligence & Valuation Dossier')
    doc.add_paragraph(f'News window: {start} to {end} | Generated: {metadata["generated_at"][:10]}')
    c = valuation.get('currency', 'N/A')
    table(doc, ['Provider identity', 'Latest close / date', 'Market cap', 'Forward P/E'], [[metadata.get('provider_name', 'Unverified'), f'{money(valuation.get("price"), c)} / {valuation.get("price_date", "N/A")}', money(valuation.get('market_cap'), c), multiple(valuation.get('forward_pe'))]])
    heading(doc, 'Executive perspective')
    doc.add_paragraph(summary['overview'])
    for i, theme in enumerate(summary['themes'], 1):
        heading(doc, f'{i}. {theme["heading"]}')
        refs = ', '.join(f'[{n}]' for n in theme.get('source_ids', []))
        doc.add_paragraph(theme['text'] + (' ' + refs if refs else ''))
    heading(doc, 'Risks & upcoming catalysts')
    doc.add_paragraph('Risks: ' + summary['risks'])
    doc.add_paragraph('Watch: ' + summary['catalysts'])
    doc.add_page_break()
    heading(doc, '4. Valuation snapshot & underwriting framework')
    metrics = [('Enterprise value', money(valuation.get('enterprise_value'), c)), ('Revenue (TTM)', money(valuation.get('revenue_ttm'), valuation.get('financial_currency', 'Currency unconfirmed'))), ('Shares outstanding', money(valuation.get('shares_outstanding'), 'shares')), ('Trailing / forward P/E', multiple(valuation.get('trailing_pe')) + ' / ' + multiple(valuation.get('forward_pe'))), ('Price / sales (TTM)', multiple(valuation.get('price_sales'))), ('EV / revenue; EV / EBITDA', multiple(valuation.get('ev_revenue')) + '; ' + multiple(valuation.get('ev_ebitda')))]
    table(doc, ['Metric', 'Provider-reported value'], metrics)
    doc.add_paragraph(summary['valuation_commentary'])
    rows, total = sotp or ([], None)
    if rows:
        heading(doc, 'Sum-of-the-parts — analyst assumptions, not provider estimates')
        table(doc, ['Segment / driver', 'Assumed multiple', 'Implied EV'], [[r['name']+' / '+r['driver'], multiple(r['multiple']), money(r['ev'], c)] for r in rows])
        doc.add_paragraph(f'Total EV: {money(total["ev"], c)}; less net debt: {money(total["net_debt"], c)}; equity: {money(total["equity"], c)}; implied diluted value/share: {money(total["per_share"], c)}. Assumptions: {total["as_of"]}; source: {total["source"]}.')
    else:
        doc.add_paragraph('SOTP / target price: not calculated. Segment forecasts, valuation multiples, net debt and diluted shares require sourced analyst assumptions; provider multiples do not establish fair value.')
    heading(doc, 'Source notes & coverage')
    doc.add_paragraph(f'{metadata["article_count"]} eligible unique articles; {summary.get("selected_article_count", 0)} included in synthesis. Financial period: {valuation.get("financial_period", "N/A")}. Fundamentals fetched: {metadata.get("fundamentals_retrieved", "N/A")}. Price fetched: {metadata.get("price_retrieved", "N/A")}. Multiples are provider snapshots and may differ in timing from the last close.')
    warnings = metadata.get('warnings', [])
    if warnings:
        doc.add_paragraph('Limitations: ' + ' '.join(warnings))
    referenced = sorted({i for t in summary['themes'] for i in t.get('source_ids', [])})
    for i in referenced:
        a = summary['sources'][i-1]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        hyperlink(p, f'[{i}] {str(a.get("date", ""))[:10]} — {str(a.get("title", ""))[:110]}', a.get('link', ''))
    doc.add_paragraph('Full article links, cached inputs, metrics and synthesis are saved alongside this document. N/A = unavailable; N/M = non-positive multiple. News reports and AI analysis require human review.')
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def render_pdf(docx_path, output_dir):
    """Return a rendered page count, not the advisory DOCX app.xml count."""
    import shutil
    if not shutil.which('libreoffice'):
        raise RuntimeError('Two-page validation requires LibreOffice (libreoffice command) and pypdf')
    from pypdf import PdfReader
    with tempfile.TemporaryDirectory() as profile:
        try:
            proc = subprocess.run(['libreoffice', '-env:UserInstallation='+Path(profile).as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(output_dir), str(docx_path)], capture_output=True, timeout=90)
        except subprocess.TimeoutExpired:
            raise RuntimeError('LibreOffice conversion timed out') from None
        if proc.returncode:
            raise RuntimeError('LibreOffice PDF conversion failed')
    pdf = Path(output_dir) / (Path(docx_path).stem+'.pdf')
    if not pdf.exists():
        raise RuntimeError('LibreOffice produced no PDF')
    return len(PdfReader(pdf).pages)
