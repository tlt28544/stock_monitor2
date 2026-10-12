import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from docx import Document
from stock_monitor.api import APIError, EODHD, merge_news, tavily_news, openai_summary
from stock_monitor.data import three_months_before, valuation_snapshot, validate_identity, sotp_rows
from stock_monitor.__main__ import main, demo_inputs, demo_summary
from stock_monitor.report import build_report


class WorkflowTests(unittest.TestCase):
    def test_cache_omits_credentials_and_preserves_timestamp(self):
        with tempfile.TemporaryDirectory() as tmp:
            c=EODHD('synthetic-secret-for-test', tmp)
            with patch('stock_monitor.api.request_json',return_value={'General': {'Code':'AAPL'}}) as request:
                c.get('fundamentals/AAPL.US')
                first=c.observations['fundamentals/AAPL.US']
                c.get('fundamentals/AAPL.US')
                self.assertEqual(request.call_count,1)
                self.assertEqual(first,c.observations['fundamentals/AAPL.US'])
            self.assertNotIn('synthetic-secret-for-test',next(Path(tmp).glob('*.json')).read_text())

    def test_failed_companies_do_not_stop_batch(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict('os.environ',{'EODHD_API_TOKEN':'test','STOCK_OPENAI_API_KEY':'test','TAVILY_API_KEY':'test'}), patch('stock_monitor.__main__.EODHD') as client, patch('stock_monitor.__main__.tavily_news',return_value=[]):
            client.return_value.get.side_effect=APIError('HTTP 403')
            client.return_value.observations={}
            self.assertEqual(main(['--no-render','--tickers','AAPL,SPCX','--output',tmp,'--as-of','2026-10-12']),1)
            manifest=json.loads((Path(tmp)/'2026-10-12/manifest.json').read_text())
            self.assertEqual([c['status'] for c in manifest['companies']],['failed','failed'])
            self.assertEqual(len(list(Path(tmp).rglob('*_failure.json'))),2)
            self.assertFalse(list(Path(tmp).rglob('*.docx')))

    def test_calendar_window(self):
        self.assertEqual(three_months_before(date(2024,5,31)), date(2024,2,29))
        self.assertEqual(three_months_before(date(2026,1,31)), date(2025,10,31))

    def test_company_list(self):
        companies = json.loads(Path('companies.json').read_text())
        self.assertEqual(len(companies),40)
        self.assertEqual(len({c['ticker'] for c in companies}),40)
        self.assertFalse(next(c for c in companies if c['ticker']=='SPCX')['verified'])

    def test_pagination_dedup_and_dates(self):
        with tempfile.TemporaryDirectory() as tmp:
            c=EODHD('test', tmp)
            first=[{'date':'2026-10-01','title':str(i),'link':f'https://x/{i}'} for i in range(100)]
            with patch.object(c,'get',side_effect=[first,[first[0],{'date':'2020-01-01','title':'old','link':'https://old'}]]):
                news,truncated=c.news('AAPL.US',date(2026,7,12),date(2026,10,12))
            self.assertEqual(len(news),100); self.assertFalse(truncated)
            with patch.object(c,'get',return_value=first):
                self.assertTrue(c.news('AAPL.US',date(2026,7,12),date(2026,10,12),1)[1])

    def test_tavily_dates_and_merge(self):
        fixture={'results':[{'published_date':'Mon, 12 Oct 2026 09:00:00 GMT','title':'News','url':'https://x/a?utm=1','content':'Evidence'}, {'published_date':None,'title':'Undated'}, {'published_date':'2020-01-01','title':'Old'}]}
        with patch('stock_monitor.api.request_json',return_value=fixture):
            news=tavily_news('test',{'name':'Apple','ticker':'AAPL'},date(2026,7,12),date(2026,10,12))
        self.assertEqual(len(news),1)
        self.assertEqual(len(merge_news(news,[{**news[0],'link':'https://x/a?utm=2'}])),1)

    def test_sotp_ev_to_equity_bridge(self):
        rows,total=sotp_rows({'currency':'USD','as_of':'2026-10-12','source':'test','net_debt':20,'diluted_shares':10,'segments':[{'name':'A','metric':'Revenue','metric_value':100,'multiple':3}]},{'currency':'USD'})
        self.assertEqual(total['ev'],300)
        self.assertEqual(total['equity'],280)
        self.assertEqual(total['per_share'],28)
        with self.assertRaises(ValueError):
            sotp_rows({'currency':'EUR'},{'currency':'USD'})

    def test_missing_and_negative_metrics(self):
        v=valuation_snapshot({'General':{},'Valuation':{'TrailingPE':-3}},[{'date':'2026-10-09','close':100,'adjusted_close':50}])
        self.assertEqual(v['price'],100)
        self.assertEqual(v['trailing_pe'],-3)
        self.assertIsNone(v['market_cap'])
        with self.assertRaises(ValueError):
            validate_identity({'symbol':'AAPL.US'},{'General':{'Code':'OTHER','Type':'ETF'}})

    def test_ai_citation_rejection(self):
        articles=[{'date':'2026-10-01','title':'Evidence'}]
        summary=demo_summary(articles)
        summary['overview']='Overview'
        response={'choices':[{'message':{'content':json.dumps(summary)}}]}
        with patch('stock_monitor.api.request_json',return_value=response):
            with self.assertRaises(APIError):
                openai_summary('test','model',{},articles,{})

    def test_docx_structure_and_demo_batch(self):
        with tempfile.TemporaryDirectory() as tmp:
            code=main(['--demo','--no-render','--tickers','AAPL,SPCX','--output',tmp,'--as-of','2026-10-12'])
            self.assertEqual(code,0)
            manifest=json.loads((Path(tmp)/'2026-10-12_DEMO/manifest.json').read_text())
            self.assertEqual(len(manifest['companies']),2)
            for item in manifest['companies']:
                doc=Document(item['document'])
                xml=doc._element.xml
                self.assertEqual(xml.count('w:type="page"'),1)
                self.assertEqual(len(doc.tables),2)
                self.assertIn('SYNTHETIC', '\n'.join(p.text for p in doc.paragraphs))
                self.assertIsNone(item['pages'])

    def test_fallback_trigger_and_continuation(self):
        fundamentals={'General':{'Code':'AAPL','Name':'Apple Inc.','Type':'Common Stock'}}
        articles=[{'date':'2026-10-01','title':'Evidence','link':'https://example.com/1'}]
        fake_summary=demo_summary(articles)
        fake_summary['themes']=fake_summary['themes'][:1]
        with tempfile.TemporaryDirectory() as tmp, patch.dict('os.environ',{'EODHD_API_TOKEN':'test','STOCK_OPENAI_API_KEY':'test','TAVILY_API_KEY':'test'}), patch('stock_monitor.__main__.EODHD') as client, patch('stock_monitor.__main__.tavily_news',return_value=articles) as fallback, patch('stock_monitor.__main__.openai_summary',return_value=fake_summary):
            client.return_value.get.side_effect=[fundamentals,[]]
            client.return_value.news.return_value=([],False)
            client.return_value.observations={}
            self.assertEqual(main(['--no-render','--tickers','AAPL','--output',tmp,'--as-of','2026-10-12']),0)
            fallback.assert_called_once()
            manifest=json.loads((Path(tmp)/'2026-10-12/manifest.json').read_text())
            self.assertEqual(manifest['companies'][0]['status'],'draft')


if __name__=='__main__':
    unittest.main()
